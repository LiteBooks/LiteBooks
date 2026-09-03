from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db.models import Q, Sum
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import (
    AccountForm,
    AccountingPeriodForm,
    BusinessDocumentForm,
    ContactForm,
    EntryHeaderForm,
    JournalLineFormSet,
    OwnerActivityForm,
    PaymentForm,
    QuickTransactionForm,
    UserCreateForm,
    UserUpdateForm,
)
from .models import (
    Account,
    AccountingPeriod,
    Attachment,
    AuditEvent,
    BusinessDocument,
    Contact,
    JournalEntry,
    JournalLine,
    OwnerActivity,
    Payment,
    UserProfile,
)
from .permissions import role_required
from .reporting import account_rows, aging, balance_sheet, income_statement, owner_balances, trial_balance
from .services import create_document, create_entry, create_owner_activity, create_payment, record_audit, update_entry


def parse_date(value, fallback=None):
    try:
        return date.fromisoformat(value) if value else fallback
    except ValueError:
        return fallback


@login_required
def dashboard(request):
    today = date.today()
    bs = balance_sheet(today)
    pnl = income_statement(date(today.year, 1, 1), today)
    cash = sum((row["balance"] for row in account_rows(end=today) if row["account"].subtype in [Account.Subtype.BANK, Account.Subtype.CASH]), Decimal("0.00"))
    recent = JournalEntry.objects.filter(status=JournalEntry.Status.POSTED).select_related("contact").prefetch_related("lines")[:8]
    open_invoices = BusinessDocument.objects.filter(kind=BusinessDocument.Kind.INVOICE).exclude(status__in=[BusinessDocument.Status.PAID, BusinessDocument.Status.VOID])
    open_bills = BusinessDocument.objects.filter(kind=BusinessDocument.Kind.BILL).exclude(status__in=[BusinessDocument.Status.PAID, BusinessDocument.Status.VOID])
    return render(request, "ledger/dashboard.html", {
        "cash": cash,
        "receivables": sum((item.balance_due for item in open_invoices), Decimal("0.00")),
        "payables": sum((item.balance_due for item in open_bills), Decimal("0.00")),
        "net_income": pnl["net_income"],
        "recent_entries": recent,
    })


@login_required
def transaction_list(request):
    entries = JournalEntry.objects.select_related("contact", "linked_entry").prefetch_related("lines__account").all()
    query = request.GET.get("q", "").strip()
    account_id = request.GET.get("account", "")
    start = parse_date(request.GET.get("start"))
    end = parse_date(request.GET.get("end"))
    source = request.GET.get("source", "")
    if query:
        entries = entries.filter(Q(number__icontains=query) | Q(description__icontains=query) | Q(memo__icontains=query) | Q(contact__name__icontains=query)).distinct()
    if account_id:
        entries = entries.filter(lines__account_id=account_id).distinct()
    if start:
        entries = entries.filter(date__gte=start)
    if end:
        entries = entries.filter(date__lte=end)
    if source:
        entries = entries.filter(source=source)
    return render(request, "ledger/transaction_list.html", {
        "entries": entries,
        "accounts": Account.objects.filter(is_active=True),
        "source_choices": JournalEntry.Source.choices,
        "filters": {"q": query, "account": account_id, "start": request.GET.get("start", ""), "end": request.GET.get("end", ""), "source": source},
    })


@login_required
def transaction_detail(request, pk):
    entry = get_object_or_404(JournalEntry.objects.select_related("contact", "linked_entry", "created_by", "posted_by").prefetch_related("lines__account", "lines__owner", "attachments"), pk=pk)
    audits = AuditEvent.objects.filter(object_type="ledger.JournalEntry", object_id=str(pk))
    return render(request, "ledger/transaction_detail.html", {"entry": entry, "audits": audits})


@role_required(UserProfile.Role.BOOKKEEPER)
def transaction_create(request):
    initial = {}
    account = request.GET.get("account")
    side = request.GET.get("side")
    if account and side in {"debit", "credit"}:
        initial[f"{side}_account"] = account
    form = QuickTransactionForm(request.POST or None, request.FILES or None, initial=initial)
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        try:
            entry = create_entry(
                entry_date=data["date"],
                description=data["description"],
                memo=data["memo"],
                contact=data["contact"],
                linked_entry=data["linked_entry"],
                attachment=data["attachment"],
                user=request.user,
                lines=[
                    {"account": data["debit_account"], "debit": data["amount"], "credit": 0, "owner": None},
                    {"account": data["credit_account"], "debit": 0, "credit": data["amount"], "owner": None},
                ],
            )
        except ValidationError as error:
            form.add_error(None, error)
        else:
            messages.success(request, f"Transaction {entry.number} posted.")
            return redirect("transaction-detail", pk=entry.pk)
    return render(request, "ledger/transaction_form.html", {"form": form, "title": "New transaction", "quick": True})


def formset_lines(formset):
    return [
        {
            "account": row["account"],
            "description": row.get("description", ""),
            "debit": row.get("debit") or 0,
            "credit": row.get("credit") or 0,
            "owner": row.get("owner"),
        }
        for row in formset.cleaned_data
        if row and not row.get("DELETE") and row.get("account")
    ]


@role_required(UserProfile.Role.BOOKKEEPER)
def transaction_split_create(request):
    entry = JournalEntry(created_by=request.user)
    form = EntryHeaderForm(request.POST or None, request.FILES or None, instance=entry)
    formset = JournalLineFormSet(request.POST or None, instance=entry)
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        data = form.cleaned_data
        try:
            saved = create_entry(
                entry_date=data["date"], description=data["description"], memo=data["memo"], contact=data["contact"],
                linked_entry=data["linked_entry"], attachment=data["attachment"], user=request.user, lines=formset_lines(formset),
            )
        except ValidationError as error:
            form.add_error(None, error)
        else:
            messages.success(request, f"Transaction {saved.number} posted.")
            return redirect("transaction-detail", pk=saved.pk)
    return render(request, "ledger/transaction_form.html", {"form": form, "formset": formset, "title": "New split transaction"})


@role_required(UserProfile.Role.BOOKKEEPER)
def transaction_edit(request, pk):
    entry = get_object_or_404(JournalEntry.objects.prefetch_related("lines"), pk=pk)
    form = EntryHeaderForm(request.POST or None, request.FILES or None, instance=entry)
    formset = JournalLineFormSet(request.POST or None, instance=entry, extra=0)
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        data = form.cleaned_data
        try:
            update_entry(
                entry=entry, entry_date=data["date"], description=data["description"], memo=data["memo"], contact=data["contact"],
                linked_entry=data["linked_entry"], attachment=data["attachment"], user=request.user, lines=formset_lines(formset),
            )
        except ValidationError as error:
            form.add_error(None, error)
        else:
            messages.success(request, f"Transaction {entry.number} updated. Version {entry.version} is now posted.")
            return redirect("transaction-detail", pk=entry.pk)
    return render(request, "ledger/transaction_form.html", {"form": form, "formset": formset, "title": f"Edit {entry.number}", "entry": entry})


@login_required
def attachment_download(request, pk):
    attachment = get_object_or_404(Attachment, pk=pk)
    try:
        return FileResponse(attachment.file.open("rb"), as_attachment=True, filename=attachment.original_name)
    except FileNotFoundError as exc:
        raise Http404("Attachment file is missing.") from exc


@login_required
def account_list(request):
    rows = []
    for account in Account.objects.annotate(debits=Sum("journal_lines__debit", filter=Q(journal_lines__entry__status="posted")), credits=Sum("journal_lines__credit", filter=Q(journal_lines__entry__status="posted"))):
        raw = (account.debits or 0) - (account.credits or 0)
        account.display_balance = raw if account.normal_balance == "debit" else -raw
        rows.append(account)
    return render(request, "ledger/account_list.html", {"accounts": rows})


@login_required
def account_detail(request, pk):
    account = get_object_or_404(Account, pk=pk)
    lines = JournalLine.objects.filter(account=account, entry__status="posted").select_related("entry", "owner")
    running = Decimal("0.00")
    ledger_rows = []
    for line in lines.order_by("entry__date", "entry__number", "sort_order"):
        movement = line.debit - line.credit
        running += movement if account.normal_balance == "debit" else -movement
        ledger_rows.append({"line": line, "balance": running})
    return render(request, "ledger/account_detail.html", {"account": account, "ledger_rows": ledger_rows, "balance": running})


@role_required(UserProfile.Role.ADMIN)
def account_create(request):
    form = AccountForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        account = form.save()
        record_audit(request.user, "created", account, after={"code": account.code, "name": account.name, "type": account.type, "subtype": account.subtype})
        messages.success(request, f"Account {account} created.")
        return redirect("account-detail", pk=account.pk)
    return render(request, "ledger/model_form.html", {"form": form, "title": "New account"})


@role_required(UserProfile.Role.ADMIN)
def account_edit(request, pk):
    account = get_object_or_404(Account, pk=pk)
    before = {"code": account.code, "name": account.name, "type": account.type, "subtype": account.subtype, "description": account.description, "is_active": account.is_active}
    form = AccountForm(request.POST or None, instance=account)
    if request.method == "POST" and form.is_valid():
        account = form.save()
        record_audit(request.user, "updated", account, before=before, after={"code": account.code, "name": account.name, "type": account.type, "subtype": account.subtype, "description": account.description, "is_active": account.is_active})
        messages.success(request, "Account updated.")
        return redirect("account-detail", pk=account.pk)
    return render(request, "ledger/model_form.html", {"form": form, "title": f"Edit {account.name}"})


@login_required
def contact_list(request):
    contacts = Contact.objects.all()
    kind = request.GET.get("kind")
    if kind:
        contacts = contacts.filter(kind=kind)
    return render(request, "ledger/contact_list.html", {"contacts": contacts, "kind_choices": Contact.Kind.choices, "selected_kind": kind})


@role_required(UserProfile.Role.BOOKKEEPER)
def contact_create(request):
    form = ContactForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        contact = form.save()
        record_audit(request.user, "created", contact, after={"name": contact.name, "kind": contact.kind})
        messages.success(request, "Contact created.")
        return redirect("contact-list")
    return render(request, "ledger/model_form.html", {"form": form, "title": "New contact"})


@role_required(UserProfile.Role.BOOKKEEPER)
def contact_edit(request, pk):
    contact = get_object_or_404(Contact, pk=pk)
    before = {"name": contact.name, "kind": contact.kind, "is_active": contact.is_active}
    form = ContactForm(request.POST or None, instance=contact)
    if request.method == "POST" and form.is_valid():
        contact = form.save()
        record_audit(request.user, "updated", contact, before=before, after={"name": contact.name, "kind": contact.kind, "is_active": contact.is_active})
        messages.success(request, "Contact updated.")
        return redirect("contact-list")
    return render(request, "ledger/model_form.html", {"form": form, "title": f"Edit {contact.name}"})


@login_required
def document_list(request):
    kind = request.GET.get("kind", BusinessDocument.Kind.INVOICE)
    if kind not in BusinessDocument.Kind.values:
        kind = BusinessDocument.Kind.INVOICE
    documents = BusinessDocument.objects.filter(kind=kind).select_related("contact")
    return render(request, "ledger/document_list.html", {"documents": documents, "kind": kind, "kind_label": BusinessDocument.Kind(kind).label})


@login_required
def document_detail(request, pk):
    document = get_object_or_404(BusinessDocument.objects.select_related("contact", "journal_entry", "control_account", "category_account").prefetch_related("payments__journal_entry"), pk=pk)
    return render(request, "ledger/document_detail.html", {"document": document})


@role_required(UserProfile.Role.BOOKKEEPER)
def document_create(request, kind):
    if kind not in BusinessDocument.Kind.values:
        raise Http404
    form = BusinessDocumentForm(request.POST or None, request.FILES or None, kind=kind)
    if request.method == "POST" and form.is_valid():
        try:
            document = create_document(kind=kind, user=request.user, attachment=form.cleaned_data.pop("attachment"), **form.cleaned_data)
        except ValidationError as error:
            form.add_error(None, error)
        else:
            messages.success(request, f"{document.get_kind_display()} {document.number} posted.")
            return redirect("document-detail", pk=document.pk)
    return render(request, "ledger/model_form.html", {"form": form, "title": f"New {BusinessDocument.Kind(kind).label.lower()}", "submit_label": "Save and post"})


@role_required(UserProfile.Role.BOOKKEEPER)
def payment_create(request, pk):
    document = get_object_or_404(BusinessDocument, pk=pk)
    form = PaymentForm(request.POST or None, request.FILES or None, document=document)
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        try:
            create_payment(document=document, payment_date=data["date"], amount=data["amount"], cash_account=data["cash_account"], user=request.user, memo=data["memo"], attachment=data["attachment"])
        except ValidationError as error:
            form.add_error(None, error)
        else:
            messages.success(request, "Payment posted and linked to the original transaction.")
            return redirect("document-detail", pk=document.pk)
    return render(request, "ledger/model_form.html", {"form": form, "title": f"Pay {document}", "submit_label": "Post payment"})


@login_required
def owner_list(request):
    return render(request, "ledger/owner_list.html", {"rows": owner_balances(), "activities": OwnerActivity.objects.select_related("owner", "journal_entry")[:20]})


@role_required(UserProfile.Role.BOOKKEEPER)
def owner_activity_create(request):
    form = OwnerActivityForm(request.POST or None, request.FILES or None)
    if request.method == "POST" and form.is_valid():
        data = form.cleaned_data
        try:
            activity = create_owner_activity(owner=data["owner"], kind=data["kind"], activity_date=data["date"], amount=data["amount"], cash_account=data["cash_account"], offset_account=data["offset_account"], user=request.user, memo=data["memo"], attachment=data["attachment"])
        except ValidationError as error:
            form.add_error(None, error)
        else:
            messages.success(request, "Owner activity posted.")
            return redirect("transaction-detail", pk=activity.journal_entry_id)
    return render(request, "ledger/model_form.html", {"form": form, "title": "New owner activity", "submit_label": "Post owner activity"})


@login_required
def reports(request):
    return render(request, "ledger/reports.html")


@login_required
def report_view(request, report):
    as_of = parse_date(request.GET.get("as_of"), date.today())
    start = parse_date(request.GET.get("start"), date(as_of.year, 1, 1))
    end = parse_date(request.GET.get("end"), as_of)
    context = {"report": report, "as_of": as_of, "start": start, "end": end}
    if report == "trial-balance":
        context["rows"], context["total_debits"], context["total_credits"] = trial_balance(as_of)
        context["title"] = "Trial balance"
    elif report == "balance-sheet":
        context["groups"] = balance_sheet(as_of)
        context["title"] = "Balance sheet"
    elif report == "income-statement":
        context.update(income_statement(start, end))
        context["title"] = "Profit and loss"
    elif report in {"receivables-aging", "payables-aging"}:
        kind = BusinessDocument.Kind.INVOICE if report == "receivables-aging" else BusinessDocument.Kind.BILL
        context["rows"], context["buckets"] = aging(kind, as_of)
        context["title"] = "Accounts receivable aging" if kind == BusinessDocument.Kind.INVOICE else "Accounts payable aging"
    elif report == "owner-balances":
        context["rows"] = owner_balances(as_of)
        context["title"] = "Owner balances"
    elif report == "general-ledger":
        context["lines"] = JournalLine.objects.filter(entry__status="posted", entry__date__range=(start, end)).select_related("entry", "account", "owner").order_by("entry__date", "entry__number", "sort_order")
        context["title"] = "General ledger"
    else:
        raise Http404
    return render(request, "ledger/report_view.html", context)


@login_required
def export_workbook(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(title, headers, rows):
        sheet = workbook.create_sheet(title)
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243746")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)

    add_sheet("Chart of Accounts", ["Code", "Name", "Type", "Subtype", "Description", "Active"], ((a.code, a.name, a.type, a.subtype, a.description, a.is_active) for a in Account.objects.all()))
    add_sheet("Transactions", ["Number", "Date", "Description", "Source", "Contact", "Linked transaction", "Memo", "Status", "Version"], ((e.number, e.date, e.description, e.source, e.contact.name if e.contact else "", e.linked_entry.number if e.linked_entry else "", e.memo, e.status, e.version) for e in JournalEntry.objects.select_related("contact", "linked_entry")))
    add_sheet("Journal Lines", ["Transaction", "Date", "Account code", "Account name", "Description", "Debit", "Credit", "Owner"], ((l.entry.number, l.entry.date, l.account.code, l.account.name, l.description, l.debit, l.credit, l.owner.name if l.owner else "") for l in JournalLine.objects.select_related("entry", "account", "owner")))
    add_sheet("Contacts", ["Name", "Type", "Email", "Phone", "Address", "Notes", "Active"], ((c.name, c.kind, c.email, c.phone, c.address, c.notes, c.is_active) for c in Contact.objects.all()))
    add_sheet("Invoices and Bills", ["Type", "Number", "Contact", "Issue date", "Due date", "Description", "Amount", "Paid", "Balance", "Status"], ((d.kind, d.number, d.contact.name, d.issue_date, d.due_date, d.description, d.amount, d.amount_paid, d.balance_due, d.status) for d in BusinessDocument.objects.select_related("contact")))
    add_sheet("Payments", ["Document type", "Document number", "Contact", "Date", "Bank or cash account", "Amount", "Transaction"], ((p.document.kind, p.document.number, p.document.contact.name, p.date, p.cash_account.code, p.amount, p.journal_entry.number) for p in Payment.objects.select_related("document__contact", "cash_account", "journal_entry")))
    add_sheet("Owner Balances", ["Owner", "Contributions", "Draws", "Amount owed"], ((r["owner"].name, r["contributions"], r["draws"], r["owed"]) for r in owner_balances()))
    tb_rows, tb_debits, tb_credits = trial_balance(date.today())
    add_sheet("Trial Balance", ["Account code", "Account name", "Debit balance", "Credit balance"], ((r["account"].code, r["account"].name, r["debit_balance"], r["credit_balance"]) for r in tb_rows if r["debit_balance"] or r["credit_balance"]))
    pnl = income_statement(date(date.today().year, 1, 1), date.today())
    add_sheet("Profit and Loss", ["Section", "Account code", "Account name", "Amount"], ((section, row["account"].code, row["account"].name, row["balance"]) for section, rows in [("Revenue", pnl["revenue"]), ("Expense", pnl["expenses"])] for row in rows))
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="litebooks-export-{date.today().isoformat()}.xlsx"'
    return response


@role_required(UserProfile.Role.ADMIN)
def period_list(request):
    form = AccountingPeriodForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        period, created = AccountingPeriod.objects.get_or_create(year=form.cleaned_data["year"], month=form.cleaned_data["month"])
        messages.success(request, f"{period} {'created' if created else 'already exists'}.")
        return redirect("period-list")
    return render(request, "ledger/period_list.html", {"periods": AccountingPeriod.objects.select_related("closed_by"), "form": form})


@role_required(UserProfile.Role.ADMIN)
@require_POST
def period_toggle(request, pk):
    period = get_object_or_404(AccountingPeriod, pk=pk)
    if period.is_closed:
        period.closed_at = None
        period.closed_by = None
        action = "reopened"
    else:
        period.closed_at = timezone.now()
        period.closed_by = request.user
        action = "closed"
    period.save(update_fields=["closed_at", "closed_by"])
    record_audit(request.user, action, period, after={"year": period.year, "month": period.month, "is_closed": period.is_closed})
    messages.success(request, f"{period} {action}.")
    return redirect("period-list")


@role_required(UserProfile.Role.ADMIN)
def user_list(request):
    form = UserCreateForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        record_audit(request.user, "created", user.profile, after={"username": user.username, "role": user.profile.role, "is_active": user.is_active})
        messages.success(request, f"User {user.username} created.")
        return redirect("user-list")
    return render(request, "ledger/user_list.html", {"users": get_user_model().objects.select_related("profile"), "form": form})


@role_required(UserProfile.Role.ADMIN)
def user_edit(request, pk):
    user = get_object_or_404(get_user_model().objects.select_related("profile"), pk=pk)
    before = {"role": user.profile.role, "is_active": user.is_active}
    form = UserUpdateForm(request.POST or None, instance=user)
    if request.method == "POST" and form.is_valid():
        if user == request.user and not form.cleaned_data["is_active"]:
            form.add_error("is_active", "You cannot deactivate your own login.")
        else:
            user = form.save()
            record_audit(request.user, "updated", user.profile, before=before, after={"username": user.username, "role": user.profile.role, "is_active": user.is_active})
            messages.success(request, f"User {user.username} updated.")
            return redirect("user-list")
    return render(request, "ledger/model_form.html", {"form": form, "title": f"Edit {user.username}", "submit_label": "Save user"})


@login_required
def audit_list(request):
    return render(request, "ledger/audit_list.html", {"events": AuditEvent.objects.select_related("actor")[:500]})

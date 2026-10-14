from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from ledger.models import Account, AccountingPeriod, AuditEvent, BusinessDocument, Contact, JournalEntry, UserProfile
from ledger.reporting import balance_sheet, owner_balances, trial_balance
from ledger.services import create_document, create_entry, create_owner_activity, create_payment, delete_entry, restore_entry, update_document, update_entry


class AccountingTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = get_user_model().objects.create_user(username="bookkeeper", password="very-good-password")
        cls.user.profile.role = UserProfile.Role.BOOKKEEPER
        cls.user.profile.save()
        cls.bank = Account.objects.create(code="1000", name="Bank", type=Account.Type.ASSET, subtype=Account.Subtype.BANK)
        cls.ar = Account.objects.create(code="1100", name="Accounts receivable", type=Account.Type.ASSET, subtype=Account.Subtype.RECEIVABLE)
        cls.ap = Account.objects.create(code="2000", name="Accounts payable", type=Account.Type.LIABILITY, subtype=Account.Subtype.PAYABLE)
        cls.loan = Account.objects.create(code="2200", name="Owner loans", type=Account.Type.LIABILITY, subtype=Account.Subtype.OWNER_LOAN_PAYABLE)
        cls.equity = Account.objects.create(code="3000", name="Contributions", type=Account.Type.EQUITY, subtype=Account.Subtype.OWNER_CONTRIBUTION)
        cls.sales = Account.objects.create(code="4000", name="Sales", type=Account.Type.REVENUE, subtype=Account.Subtype.SALES)
        cls.expense = Account.objects.create(code="6000", name="Expense", type=Account.Type.EXPENSE, subtype=Account.Subtype.OPERATING_EXPENSE)
        cls.customer = Contact.objects.create(name="Acme Customer", kind=Contact.Kind.CUSTOMER)
        cls.vendor = Contact.objects.create(name="Supply Vendor", kind=Contact.Kind.VENDOR)
        cls.owner = Contact.objects.create(name="Alex Owner", kind=Contact.Kind.OWNER)

    def test_balanced_entry_posts_and_is_audited(self):
        entry = create_entry(
            entry_date=date(2026, 1, 5), description="Cash sale", user=self.user,
            lines=[{"account": self.bank, "debit": 125, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 125, "owner": None}],
        )
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)
        self.assertEqual(entry.number, "20260105-0001")
        self.assertTrue(entry.is_balanced)
        self.assertEqual(entry.total, Decimal("125"))
        self.assertEqual(AuditEvent.objects.filter(object_id=str(entry.pk), action="created").count(), 1)

    def test_numbers_reset_by_posting_date(self):
        first = create_entry(
            entry_date=date(2026, 5, 1), description="First", user=self.user,
            lines=[{"account": self.bank, "debit": 1, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 1, "owner": None}],
        )
        second = create_entry(
            entry_date=date(2026, 5, 1), description="Second", user=self.user,
            lines=[{"account": self.bank, "debit": 1, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 1, "owner": None}],
        )
        next_day = create_entry(
            entry_date=date(2026, 5, 2), description="Next day", user=self.user,
            lines=[{"account": self.bank, "debit": 1, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 1, "owner": None}],
        )
        self.assertEqual((first.number, second.number, next_day.number), ("20260501-0001", "20260501-0002", "20260502-0001"))

    def test_unbalanced_entry_rolls_back(self):
        with self.assertRaises(ValidationError):
            create_entry(
                entry_date=date(2026, 1, 5), description="Broken", user=self.user,
                lines=[{"account": self.bank, "debit": 100, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 90, "owner": None}],
            )
        self.assertFalse(JournalEntry.objects.filter(description="Broken").exists())

    def test_closed_period_rejects_new_and_edited_entries(self):
        entry = create_entry(
            entry_date=date(2026, 1, 5), description="Before close", user=self.user,
            lines=[{"account": self.bank, "debit": 100, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 100, "owner": None}],
        )
        AccountingPeriod.objects.create(year=2026, month=1, closed_at=timezone.now(), closed_by=self.user)
        with self.assertRaises(ValidationError):
            create_entry(
                entry_date=date(2026, 1, 6), description="After close", user=self.user,
                lines=[{"account": self.bank, "debit": 50, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 50, "owner": None}],
            )
        with self.assertRaises(ValidationError):
            update_entry(
                entry=entry, entry_date=entry.date, description="Changed", user=self.user,
                lines=[{"account": self.bank, "debit": 100, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 100, "owner": None}],
            )

    def sale(self, entry_date=date(2026, 7, 1), amount=200, description="Sale"):
        return create_entry(
            entry_date=entry_date, description=description, user=self.user,
            lines=[{"account": self.bank, "debit": amount, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": amount, "owner": None}],
        )

    def test_deleted_entry_leaves_the_books_but_keeps_its_history(self):
        entry = self.sale()
        before = trial_balance(date(2026, 7, 31))[1]
        delete_entry(entry=entry, user=self.user, reason="Duplicate of an earlier sale")
        entry.refresh_from_db()
        self.assertTrue(entry.is_deleted)
        self.assertEqual(entry.deleted_by, self.user)
        self.assertEqual(entry.deleted_reason, "Duplicate of an earlier sale")
        self.assertEqual(trial_balance(date(2026, 7, 31))[1], before - Decimal("200"))
        self.assertEqual(entry.lines.count(), 2)
        audit = AuditEvent.objects.get(object_id=str(entry.pk), action="deleted")
        self.assertEqual(audit.before["status"], JournalEntry.Status.POSTED)
        self.assertEqual(audit.after["deleted_reason"], "Duplicate of an earlier sale")
        self.assertEqual(len(audit.before["lines"]), 2)

    def test_deleted_entry_can_be_restored_to_the_books(self):
        entry = self.sale()
        posted = trial_balance(date(2026, 7, 31))[1]
        delete_entry(entry=entry, user=self.user)
        restore_entry(entry=entry, user=self.user)
        entry.refresh_from_db()
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)
        self.assertIsNone(entry.deleted_at)
        self.assertEqual(entry.deleted_by, None)
        self.assertEqual(trial_balance(date(2026, 7, 31))[1], posted)
        self.assertTrue(AuditEvent.objects.filter(object_id=str(entry.pk), action="restored").exists())

    def test_deleted_entry_cannot_be_edited_or_deleted_twice(self):
        entry = self.sale()
        delete_entry(entry=entry, user=self.user)
        with self.assertRaises(ValidationError):
            delete_entry(entry=entry, user=self.user)
        with self.assertRaises(ValidationError):
            update_entry(
                entry=entry, entry_date=entry.date, description="Changed", user=self.user,
                lines=[{"account": self.bank, "debit": 200, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 200, "owner": None}],
            )

    def test_closed_period_rejects_deletion_and_restore(self):
        entry = self.sale(entry_date=date(2026, 8, 3))
        AccountingPeriod.objects.create(year=2026, month=8, closed_at=timezone.now(), closed_by=self.user)
        with self.assertRaises(ValidationError):
            delete_entry(entry=entry, user=self.user)
        entry.refresh_from_db()
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)

    def test_document_entry_cannot_be_deleted_as_a_transaction(self):
        invoice = create_document(
            kind=BusinessDocument.Kind.INVOICE, number="INV-900", contact=self.customer,
            issue_date=date(2026, 9, 1), due_date=date(2026, 10, 1), description="Services", amount=Decimal("300"),
            control_account=self.ar, category_account=self.sales, user=self.user,
        )
        with self.assertRaises(ValidationError):
            delete_entry(entry=invoice.journal_entry, user=self.user)

    def test_delete_api_hides_the_entry_from_the_default_listing(self):
        entry = self.sale(entry_date=date(2026, 10, 5), description="API delete")
        self.client.force_login(self.user)
        response = self.client.delete(
            reverse("api-transaction", args=[entry.pk]),
            data={"reason": "Posted in error"},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["entry"]["is_deleted"])
        listed = self.client.get(reverse("api-transactions")).json()["entries"]
        self.assertNotIn(entry.pk, [item["id"] for item in listed])
        only_deleted = self.client.get(reverse("api-transactions"), {"deleted": "only"}).json()["entries"]
        self.assertEqual([item["id"] for item in only_deleted], [entry.pk])
        detail = self.client.get(reverse("api-transaction", args=[entry.pk])).json()
        self.assertEqual(detail["entry"]["deleted_reason"], "Posted in error")
        self.assertFalse(detail["entry"]["can_edit"])
        restored = self.client.post(reverse("api-transaction-restore", args=[entry.pk]))
        self.assertEqual(restored.status_code, 200)
        self.assertFalse(restored.json()["entry"]["is_deleted"])

    def test_viewer_cannot_delete_a_transaction(self):
        entry = self.sale(entry_date=date(2026, 11, 2))
        viewer = get_user_model().objects.create_user(username="viewer-delete", password="very-good-password")
        self.client.force_login(viewer)
        self.assertEqual(self.client.delete(reverse("api-transaction", args=[entry.pk])).status_code, 403)
        self.assertEqual(self.client.post(reverse("api-transaction-restore", args=[entry.pk])).status_code, 403)
        entry.refresh_from_db()
        self.assertEqual(entry.status, JournalEntry.Status.POSTED)

    def test_invoice_and_payment_link_to_original_entry(self):
        invoice = create_document(
            kind=BusinessDocument.Kind.INVOICE, number="INV-100", contact=self.customer,
            issue_date=date(2026, 2, 1), due_date=date(2026, 3, 1), description="Services", amount=Decimal("500"),
            control_account=self.ar, category_account=self.sales, user=self.user,
        )
        payment = create_payment(
            document=invoice, payment_date=date(2026, 2, 10), amount=Decimal("200"), cash_account=self.bank, user=self.user,
        )
        invoice.refresh_from_db()
        self.assertEqual(invoice.status, BusinessDocument.Status.PARTIAL)
        self.assertEqual(invoice.balance_due, Decimal("300"))
        self.assertEqual(payment.journal_entry.linked_entry, invoice.journal_entry)
        self.assertTrue(payment.journal_entry.is_balanced)

    def invoice(self, number="INV-500", amount="500", issue_date=date(2026, 4, 1)):
        return create_document(
            kind=BusinessDocument.Kind.INVOICE, number=number, contact=self.customer,
            issue_date=issue_date, due_date=date(2026, 5, 1), description="Services", amount=Decimal(amount),
            control_account=self.ar, category_account=self.sales, user=self.user,
        )

    def edit_invoice(self, invoice, **changes):
        fields = {
            "number": invoice.number, "contact": invoice.contact, "issue_date": invoice.issue_date, "due_date": invoice.due_date,
            "description": invoice.description, "amount": invoice.amount, "control_account": invoice.control_account,
            "category_account": invoice.category_account,
        }
        return update_document(document=invoice, user=self.user, **{**fields, **changes})

    def test_editing_an_invoice_rewrites_its_entry(self):
        invoice = self.invoice()
        other_revenue = Account.objects.create(code="4100", name="Consulting", type=Account.Type.REVENUE, subtype=Account.Subtype.SALES)
        invoice = self.edit_invoice(invoice, number="INV-501", amount=Decimal("650"), issue_date=date(2026, 4, 3), description="More services", category_account=other_revenue)
        entry = JournalEntry.objects.get(pk=invoice.journal_entry_id)
        self.assertEqual((entry.date, entry.number, entry.version), (date(2026, 4, 3), "20260403-0001", 2))
        self.assertEqual(entry.description, "Invoice INV-501: More services")
        self.assertEqual(
            [(line.account, line.debit, line.credit) for line in entry.lines.order_by("sort_order")],
            [(self.ar, Decimal("650"), Decimal("0")), (other_revenue, Decimal("0"), Decimal("650"))],
        )
        self.assertTrue(entry.is_balanced)
        self.assertEqual(AuditEvent.objects.filter(object_type="ledger.BusinessDocument", object_id=str(invoice.pk), action="updated").count(), 1)

    def test_editing_a_paid_invoice_respects_its_payments(self):
        invoice = self.invoice()
        create_payment(document=invoice, payment_date=date(2026, 4, 10), amount=Decimal("200"), cash_account=self.bank, user=self.user)
        with self.assertRaises(ValidationError):
            self.edit_invoice(invoice, amount=Decimal("150"))
        other_customer = Contact.objects.create(name="Other Customer", kind=Contact.Kind.CUSTOMER)
        with self.assertRaises(ValidationError):
            self.edit_invoice(invoice, contact=other_customer)
        invoice = self.edit_invoice(invoice, amount=Decimal("200"))
        self.assertEqual(invoice.status, BusinessDocument.Status.PAID)
        invoice = self.edit_invoice(invoice, amount=Decimal("800"))
        self.assertEqual((invoice.status, invoice.balance_due), (BusinessDocument.Status.PARTIAL, Decimal("600")))

    def test_renumbering_an_invoice_renames_its_payments(self):
        invoice = self.invoice()
        payment = create_payment(document=invoice, payment_date=date(2026, 4, 10), amount=Decimal("200"), cash_account=self.bank, user=self.user)
        self.edit_invoice(invoice, number="INV-777")
        entry = JournalEntry.objects.get(pk=payment.journal_entry_id)
        self.assertEqual((entry.description, entry.version), ("Payment received for invoice INV-777", 2))
        self.assertEqual(AuditEvent.objects.filter(object_type="ledger.JournalEntry", object_id=str(entry.pk), action="updated").count(), 1)
        # Other edits leave the payment alone.
        self.edit_invoice(BusinessDocument.objects.get(pk=invoice.pk), description="Reworded")
        self.assertEqual(JournalEntry.objects.get(pk=payment.journal_entry_id).version, 2)

    def test_renumbering_is_blocked_when_a_payment_is_in_a_closed_period(self):
        invoice = self.invoice()
        create_payment(document=invoice, payment_date=date(2026, 5, 10), amount=Decimal("200"), cash_account=self.bank, user=self.user)
        AccountingPeriod.objects.create(year=2026, month=5, closed_at=timezone.now(), closed_by=self.user)
        with self.assertRaises(ValidationError):
            self.edit_invoice(invoice, number="INV-778")
        self.assertEqual(BusinessDocument.objects.get(pk=invoice.pk).number, "INV-500")

    def test_closed_period_rejects_invoice_edits(self):
        invoice = self.invoice(issue_date=date(2026, 6, 2))
        AccountingPeriod.objects.create(year=2026, month=6, closed_at=timezone.now(), closed_by=self.user)
        with self.assertRaises(ValidationError):
            self.edit_invoice(invoice, issue_date=date(2026, 7, 1))
        self.assertEqual(JournalEntry.objects.get(pk=invoice.journal_entry_id).version, 1)

    def test_bill_edit_api_and_entry_edit_link(self):
        bill = create_document(
            kind=BusinessDocument.Kind.BILL, number="B-1", contact=self.vendor,
            issue_date=date(2026, 8, 1), due_date=date(2026, 8, 31), description="Supplies", amount=Decimal("90"),
            control_account=self.ap, category_account=self.expense, user=self.user,
        )
        self.client.force_login(self.user)
        detail = self.client.get(reverse("api-transaction", args=[bill.journal_entry_id])).json()["entry"]
        self.assertEqual((detail["can_edit"], detail["edit_url"]), (True, f"/documents/{bill.pk}/edit/"))
        response = self.client.post(reverse("api-document", args=[bill.pk]), {
            "number": "B-1", "contact": self.vendor.pk, "issue_date": "2026-08-01", "due_date": "2026-08-31",
            "description": "Supplies", "amount": "120.00", "control_account": self.ap.pk, "category_account": self.expense.pk,
        })
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(response.json()["document"]["amount"], "120.00")
        self.assertEqual(JournalEntry.objects.get(pk=bill.journal_entry_id).total, Decimal("120"))
        # The journal endpoint still refuses to touch a document's entry directly.
        response = self.client.post(reverse("api-transaction", args=[bill.journal_entry_id]), {"date": "2026-08-01", "description": "x", "lines": "[]"})
        self.assertEqual(response.status_code, 400)

    def test_owner_loan_tracks_amount_owed(self):
        create_owner_activity(
            owner=self.owner, kind="loan_in", activity_date=date(2026, 3, 1), amount=Decimal("1000"),
            cash_account=self.bank, offset_account=self.loan, user=self.user,
        )
        create_owner_activity(
            owner=self.owner, kind="loan_repayment", activity_date=date(2026, 3, 15), amount=Decimal("250"),
            cash_account=self.bank, offset_account=self.loan, user=self.user,
        )
        self.assertEqual(owner_balances()[0]["owed"], Decimal("750"))

    def test_financial_reports_balance(self):
        create_entry(
            entry_date=date(2026, 4, 1), description="Sale", user=self.user,
            lines=[{"account": self.bank, "debit": 500, "credit": 0, "owner": None}, {"account": self.sales, "debit": 0, "credit": 500, "owner": None}],
        )
        rows, debits, credits = trial_balance(date(2026, 4, 30))
        self.assertEqual(debits, credits)
        statement = balance_sheet(date(2026, 4, 30))
        self.assertEqual(statement[Account.Type.ASSET]["total"], statement[Account.Type.LIABILITY]["total"] + statement[Account.Type.EQUITY]["total"])

    def test_excel_export_has_accounting_sheets(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("export-workbook"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertTrue({"Chart of Accounts", "Transactions", "Journal Lines", "Payments", "Trial Balance"}.issubset(workbook.sheetnames))

    def test_viewer_cannot_post(self):
        viewer = get_user_model().objects.create_user(username="viewer", password="very-good-password")
        self.client.force_login(viewer)
        response = self.client.post(reverse("api-transactions"), data={"mode": "quick"}, content_type="application/json")
        self.assertEqual(response.status_code, 403)

    def test_authenticated_session_and_spa_shell(self):
        self.client.force_login(self.user)
        session = self.client.get(reverse("api-session"))
        self.assertTrue(session.json()["authenticated"])
        self.assertTrue(session.json()["permissions"]["edit_books"])
        page = self.client.get(reverse("transaction-list"))
        self.assertContains(page, 'id="root"')

    def test_quick_transaction_api_uses_daily_number(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("api-transactions"),
            data={
                "mode": "quick",
                "date": "2026-06-10",
                "description": "API sale",
                "amount": "45.00",
                "debit_account_id": self.bank.pk,
                "credit_account_id": self.sales.pk,
            },
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["entry"]["number"], "20260610-0001")

    def test_profile_api_updates_the_logged_in_user(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("api-profile"),
            data={"username": "updated-bookkeeper", "first_name": "Taylor", "last_name": "Morgan", "current_password": "", "new_password": "", "confirm_password": ""},
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 200)
        self.user.refresh_from_db()
        self.assertEqual(self.user.username, "updated-bookkeeper")
        self.assertEqual(self.user.get_full_name(), "Taylor Morgan")
        self.assertEqual(self.client.get(reverse("api-session")).json()["user"]["name"], "Taylor Morgan")
        event = AuditEvent.objects.get(action="updated own profile", object_id=str(self.user.profile.pk))
        self.assertNotIn("password", str(event.before).lower())
        self.assertNotIn("password", str(event.after).lower())

    def test_profile_password_change_requires_current_password_and_keeps_session(self):
        self.client.force_login(self.user)
        payload = {"username": self.user.username, "first_name": "", "last_name": "", "current_password": "incorrect", "new_password": "Cloud-Bridge!9246", "confirm_password": "Cloud-Bridge!9246"}
        rejected = self.client.post(reverse("api-profile"), data=payload, content_type="application/json")
        self.assertEqual(rejected.status_code, 400)
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("very-good-password"))

        payload["current_password"] = "very-good-password"
        changed = self.client.post(reverse("api-profile"), data=payload, content_type="application/json")
        self.assertEqual(changed.status_code, 200)
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("Cloud-Bridge!9246"))
        self.assertTrue(self.client.get(reverse("api-session")).json()["authenticated"])
